import asyncio
from datetime import date
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.auth import require_chapter_lead
from app.db import Base
from app.models import ProfessionalRole, Squad, Vendor
from app.routes import (
    create_assignment,
    create_member,
    delete_assignment,
    export_assignments,
    get_metrics_dashboard,
    list_assignment_candidates,
    list_assignments,
    update_assignment,
    update_member,
)
from app.schemas import AssignmentCreate, AssignmentUpdate, CurrentUser, MemberCreate, MemberUpdate


@pytest.fixture()
def db() -> Session:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    with session_factory() as session:
        session.add_all(
            [
                ProfessionalRole(id="role-1", name="Data Engineer"),
                ProfessionalRole(id="role-2", name="Data Analyst"),
                Vendor(id="vendor-1", name="Proveedor Uno", ruc="LEGACY-1", focal_point="Ana Contacto"),
                Vendor(id="vendor-2", name="Proveedor Dos", ruc="LEGACY-2", focal_point="Luis Contacto"),
                Squad(id="squad-1", code="DATA1", name="Squad Data"),
                Squad(id="squad-2", code="DATA2", name="Squad Analytics"),
                Squad(id="squad-inactive", code="OLD1", name="Squad Inactivo", is_active=False),
            ]
        )
        session.commit()
        yield session
    Base.metadata.drop_all(engine)


def member_payload(**overrides: object) -> dict[str, object]:
    value: dict[str, object] = {
        "dni": "12345678",
        "first_name": "Ana",
        "paternal_surname": "Pérez",
        "maternal_surname": "Gómez",
        "email": "ana@example.com",
        "birth_date": "1990-01-01",
        "employment_type": "TERCERIZADO",
        "vendor_id": "vendor-1",
        "professional_role_id": "role-1",
        "seniority": "SENIOR",
    }
    value.update(overrides)
    return value


def assignment_payload(member_id: str, **overrides: object) -> dict[str, object]:
    value: dict[str, object] = {
        "member_id": member_id,
        "assigned_squad_id": "squad-1",
        "executor_squad_id": "squad-1",
        "project_code": "PROJ-42/ETL!",
        "start_date": "2026-01-01",
        "end_date": "2026-01-31",
        "allocation_percentage": "60",
    }
    value.update(overrides)
    return value


def create_test_member(db: Session, user: CurrentUser, **overrides: object):
    return create_member(MemberCreate(**member_payload(**overrides)), user, db)


def test_assignment_schema_accepts_project_symbols_and_validates_dates() -> None:
    assignment = AssignmentCreate(
        **assignment_payload("member-1", project_code="  APP/CORE+2 - QA!  ", allocation_percentage="12.50")
    )
    assert assignment.project_code == "APP/CORE+2 - QA!"
    assert str(assignment.allocation_percentage) == "12.50"

    with pytest.raises(ValueError, match="fecha fin"):
        AssignmentCreate(**assignment_payload("member-1", start_date="2026-02-01", end_date="2026-01-31"))


def test_create_assignment_captures_vendor_and_candidate_search(db: Session) -> None:
    user = CurrentUser(user_id="lead", role="Chapter Lead")
    member = create_test_member(db, user)
    created = create_assignment(AssignmentCreate(**assignment_payload(member.id)), user, db)

    assert created.vendor_id == "vendor-1"
    assert created.vendor_name == "Proveedor Uno"
    assert created.member_full_name == "Ana Pérez Gómez"
    candidates = list_assignment_candidates("perez", user, db)
    assert [candidate.id for candidate in candidates] == [member.id]

    db.get(Vendor, "vendor-1").is_active = False
    db.commit()
    assert list_assignments("", None, None, None, user, db).items[0].vendor_name == "Proveedor Uno"

    update_member(member.id, MemberUpdate(vendor_id="vendor-2"), user, db)
    detail = list_assignments("", None, None, None, user, db).items[0]
    assert detail.vendor_id == "vendor-1"
    assert detail.vendor_name == "Proveedor Uno"


def test_capacity_allows_100_percent_and_rejects_overlapping_excess(db: Session) -> None:
    user = CurrentUser(user_id="lead", role="Chapter Lead")
    member = create_test_member(db, user)
    create_assignment(
        AssignmentCreate(**assignment_payload(member.id, allocation_percentage="60")),
        user,
        db,
    )
    second = create_assignment(
        AssignmentCreate(
            **assignment_payload(
                member.id,
                assigned_squad_id="squad-2",
                executor_squad_id="squad-2",
                project_code="PROJECT-SECOND",
                allocation_percentage="40",
                start_date="2026-01-15",
                end_date="2026-02-15",
            )
        ),
        user,
        db,
    )
    assert second.allocation_percentage == 40

    with pytest.raises(HTTPException) as error:
        create_assignment(
            AssignmentCreate(
                **assignment_payload(
                    member.id,
                    project_code="PROJECT-OVER",
                    allocation_percentage="1",
                    start_date="2026-01-20",
                    end_date="2026-01-25",
                )
            ),
            user,
            db,
        )
    assert error.value.status_code == 409
    assert "100%" in error.value.detail

    outside = create_assignment(
        AssignmentCreate(
            **assignment_payload(
                member.id,
                project_code="PROJECT-FUTURE",
                allocation_percentage="100",
                start_date="2026-03-01",
                end_date="2026-03-31",
            )
        ),
        user,
        db,
    )
    assert outside.allocation_percentage == 100


def test_assignment_filters_update_and_delete(db: Session) -> None:
    user = CurrentUser(user_id="lead", role="Chapter Lead")
    member = create_test_member(db, user)
    assignment = create_assignment(AssignmentCreate(**assignment_payload(member.id)), user, db)

    filtered = list_assignments("12345678", "vendor-1", "role-1", "squad-1", user, db)
    assert [item.id for item in filtered.items] == [assignment.id]
    updated = update_assignment(
        assignment.id,
        AssignmentUpdate(assigned_squad_id="squad-2", executor_squad_id="squad-1", project_code="NEW+PROJECT", allocation_percentage="80"),
        user,
        db,
    )
    assert updated.id == assignment.id
    assert updated.assigned_squad_id == "squad-2"
    assert updated.executor_squad_id == "squad-1"
    assert updated.vendor_id == "vendor-1"

    delete_assignment(assignment.id, user, db)
    assert list_assignments("", None, None, None, user, db).items == []


def test_assignment_list_paginates_and_sorts_stably(db: Session) -> None:
    user = CurrentUser(user_id="lead", role="Chapter Lead")
    first = create_test_member(
        db, user, dni="11111111", first_name="Zoe", email="zoe@example.com"
    )
    second = create_test_member(
        db, user, dni="22222222", first_name="Ana", email="ana2@example.com"
    )
    third = create_test_member(
        db, user, dni="33333333", first_name="Luis", email="luis2@example.com"
    )
    create_assignment(
        AssignmentCreate(**assignment_payload(first.id, project_code="PROJECT-Z")), user, db
    )
    create_assignment(
        AssignmentCreate(
            **assignment_payload(
                second.id, project_code="PROJECT-A", start_date="2026-02-01", end_date="2026-02-28"
            )
        ),
        user,
        db,
    )
    create_assignment(
        AssignmentCreate(
            **assignment_payload(
                third.id, project_code="PROJECT-L", start_date="2026-03-01", end_date="2026-03-31"
            )
        ),
        user,
        db,
    )

    first_page = list_assignments(
        "", None, None, None, user, db, page=1, page_size=2, sort_by="member_full_name", sort_direction="asc"
    )
    second_page = list_assignments(
        "", None, None, None, user, db, page=2, page_size=2, sort_by="member_full_name", sort_direction="asc"
    )
    descending = list_assignments(
        "", None, None, None, user, db, page=1, page_size=2, sort_by="project_code", sort_direction="desc"
    )

    assert (first_page.total, first_page.page, first_page.page_size, first_page.total_pages) == (3, 1, 2, 2)
    assert [item.member_full_name for item in first_page.items] == ["Ana Pérez Gómez", "Luis Pérez Gómez"]
    assert [item.member_full_name for item in second_page.items] == ["Zoe Pérez Gómez"]
    assert [item.project_code for item in descending.items] == ["PROJECT-Z", "PROJECT-L"]

    with pytest.raises(HTTPException) as invalid_page:
        list_assignments("", None, None, None, user, db, page=0)
    assert invalid_page.value.status_code == 422
    with pytest.raises(HTTPException) as invalid_sort:
        list_assignments("", None, None, None, user, db, sort_by="unsupported")  # type: ignore[arg-type]
    assert invalid_sort.value.status_code == 422


def test_assignment_date_filters_are_inclusive_and_support_partial_ranges(db: Session) -> None:
    user = CurrentUser(user_id="lead", role="Chapter Lead")
    member = create_test_member(db, user)
    assignment = create_assignment(AssignmentCreate(**assignment_payload(member.id)), user, db)

    on_start_boundary = list_assignments(
        "", None, None, None, user, db, start_date=date(2026, 1, 31), end_date=date(2026, 1, 31)
    )
    on_end_boundary = list_assignments(
        "", None, None, None, user, db, start_date=date(2026, 1, 1), end_date=date(2026, 1, 1)
    )
    before_range = list_assignments("", None, None, None, user, db, end_date=date(2025, 12, 31))
    after_range = list_assignments("", None, None, None, user, db, start_date=date(2026, 2, 1))

    assert [item.id for item in on_start_boundary.items] == [assignment.id]
    assert [item.id for item in on_end_boundary.items] == [assignment.id]
    assert before_range.items == []
    assert after_range.items == []

    with pytest.raises(HTTPException, match="filtro") as error:
        list_assignments(
            "", None, None, None, user, db, start_date=date(2026, 2, 1), end_date=date(2026, 1, 1)
        )
    assert error.value.status_code == 422


def test_inactive_squad_and_unauthorized_user_are_rejected(db: Session) -> None:
    user = CurrentUser(user_id="lead", role="Chapter Lead")
    member = create_test_member(db, user)
    with pytest.raises(HTTPException) as error:
        create_assignment(
            AssignmentCreate(**assignment_payload(member.id, assigned_squad_id="squad-inactive")),
            user,
            db,
        )
    assert error.value.status_code == 422

    with pytest.raises(HTTPException) as error:
        create_assignment(
            AssignmentCreate(**assignment_payload(member.id, executor_squad_id="squad-inactive")),
            user,
            db,
        )
    assert error.value.status_code == 422

    with pytest.raises(HTTPException) as error:
        asyncio.run(require_chapter_lead(CurrentUser(user_id="member", role="Miembro de Equipo")))
    assert error.value.status_code == 403


def test_assignment_supports_different_executor_squad(db: Session) -> None:
    user = CurrentUser(user_id="lead", role="Chapter Lead")
    member = create_test_member(db, user)
    created = create_assignment(
        AssignmentCreate(
            **assignment_payload(member.id, assigned_squad_id="squad-1", executor_squad_id="squad-2")
        ),
        user,
        db,
    )

    assert created.assigned_squad_id == "squad-1"
    assert created.assigned_squad_name == "Squad Data"
    assert created.executor_squad_id == "squad-2"
    assert created.executor_squad_name == "Squad Analytics"


def test_metrics_use_executor_squad_unique_members_and_exclude_resigned(db: Session) -> None:
    user = CurrentUser(user_id="member", role="Miembro de Equipo")
    first = create_test_member(db, user, dni="11111111", email="first@example.com")
    resigned = create_test_member(
        db, user, dni="22222222", email="resigned@example.com", professional_role_id="role-2"
    )
    create_test_member(db, user, dni="33333333", email="unassigned@example.com")

    create_assignment(
        AssignmentCreate(**assignment_payload(first.id, allocation_percentage="60")), user, db
    )
    create_assignment(
        AssignmentCreate(
            **assignment_payload(
                first.id,
                allocation_percentage="40",
                project_code="SECOND-PROJECT",
                start_date="2026-02-01",
                end_date="2026-02-28",
            )
        ),
        user,
        db,
    )
    create_assignment(
        AssignmentCreate(
            **assignment_payload(
                resigned.id,
                member_resigned=True,
                project_code="RESIGNED-PROJECT",
                allocation_percentage="100",
            )
        ),
        user,
        db,
    )
    create_assignment(
        AssignmentCreate(
            **assignment_payload(
                resigned.id,
                member_resigned=False,
                project_code="PREVIOUS-PROJECT",
                start_date="2025-10-01",
                end_date="2025-12-31",
                allocation_percentage="100",
            )
        ),
        user,
        db,
    )

    dashboard = get_metrics_dashboard(2026, 1, user, db)

    assert dashboard.period.start_date == date(2026, 1, 1)
    assert dashboard.period.end_date == date(2026, 3, 31)
    assert dashboard.kpis.total_squads == 1
    assert dashboard.kpis.assigned_members == 1
    assert dashboard.kpis.unassigned_active_members == 2
    assert dashboard.kpis.resigned_members == 1
    assert dashboard.role_distribution_by_executor_squad[0].roles[0].member_count == 1
    assert dashboard.members_by_vendor_and_role[0].vendor_name == "Proveedor Uno"
    assert {member.dni for member in dashboard.unassigned_members} == {"22222222", "33333333"}
    assert {card.project_code for card in dashboard.assignment_cards} == {"PROJ-42/ETL!", "SECOND-PROJECT"}
    assert all(card.vendor_name == "Proveedor Uno" for card in dashboard.assignment_cards)
    assert all(card.member_full_name == "Ana Pérez Gómez" for card in dashboard.assignment_cards)
    assert all(card.project_code != "RESIGNED-PROJECT" for card in dashboard.assignment_cards)
    assert {option.id for option in dashboard.assignment_filter_options.providers} == {"vendor-1", "vendor-2"}
    assert {option.id for option in dashboard.assignment_filter_options.squads} == {"squad-1", "squad-2"}



def test_metrics_assignment_cards_apply_combined_filters_and_keep_dashboard_metrics(db: Session) -> None:
    user = CurrentUser(user_id="member", role="Miembro de Equipo")
    first = create_test_member(db, user, dni="11111111", email="first@example.com")
    second = create_test_member(
        db,
        user,
        dni="22222222",
        first_name="Luis",
        email="second@example.com",
        vendor_id="vendor-2",
        professional_role_id="role-2",
    )
    create_assignment(AssignmentCreate(**assignment_payload(first.id)), user, db)
    create_assignment(
        AssignmentCreate(
            **assignment_payload(
                second.id,
                assigned_squad_id="squad-2",
                executor_squad_id="squad-1",
                project_code="PROJECT-SECOND",
            )
        ),
        user,
        db,
    )

    filtered = get_metrics_dashboard(
        2026,
        1,
        user,
        db,
        vendor_id="vendor-2",
        professional_role_id="role-2",
        assigned_squad_id=["squad-1", "squad-2"],
    )

    assert [card.project_code for card in filtered.assignment_cards] == ["PROJECT-SECOND"]
    assert filtered.assignment_cards[0].assigned_squad_name == "Squad Analytics"
    assert filtered.assignment_cards[0].executor_squad_name == "Squad Data"
    assert filtered.kpis.assigned_members == 2
    assert len(filtered.role_distribution_by_executor_squad) == 1
    assert filtered.role_distribution_by_executor_squad[0].squad_id == "squad-1"

    all_cards = get_metrics_dashboard(2026, 1, user, db, assigned_squad_id=[])
    assert {card.project_code for card in all_cards.assignment_cards} == {"PROJ-42/ETL!", "PROJECT-SECOND"}

    with pytest.raises(HTTPException) as invalid_period:
        get_metrics_dashboard(2026, 5, user, db)
    assert invalid_period.value.status_code == 422


def test_assignment_resignation_defaults_and_can_be_toggled_without_changing_dates(db: Session) -> None:
    user = CurrentUser(user_id="lead", role="Chapter Lead")
    member = create_test_member(db, user)
    assignment = create_assignment(AssignmentCreate(**assignment_payload(member.id)), user, db)

    assert assignment.member_resigned is False

    marked = update_assignment(
        assignment.id,
        AssignmentUpdate(member_resigned=True),
        user,
        db,
    )
    assert marked.member_resigned is True
    assert marked.start_date == date(2026, 1, 1)
    assert marked.end_date == date(2026, 1, 31)

    shortened = update_assignment(
        assignment.id,
        AssignmentUpdate(member_resigned=True, end_date=date(2026, 1, 15)),
        user,
        db,
    )
    assert shortened.member_resigned is True
    assert shortened.end_date == date(2026, 1, 15)

    unchanged_marker = update_assignment(
        assignment.id,
        AssignmentUpdate(project_code="PROJECT-UPDATED"),
        user,
        db,
    )
    assert unchanged_marker.member_resigned is True

    unmarked = update_assignment(assignment.id, AssignmentUpdate(member_resigned=False), user, db)
    assert unmarked.member_resigned is False
    assert unmarked.end_date == date(2026, 1, 15)


def test_assignment_resignation_filter_capacity_and_independent_replacement(db: Session) -> None:
    user = CurrentUser(user_id="lead", role="Chapter Lead")
    departing_member = create_test_member(db, user)
    replacement_member = create_test_member(
        db,
        user,
        dni="87654321",
        first_name="Luis",
        paternal_surname="Pérez",
        maternal_surname="Ruiz",
        email="luis@example.com",
    )
    departing = create_assignment(
        AssignmentCreate(
            **assignment_payload(
                departing_member.id,
                member_resigned=True,
                allocation_percentage="60",
                end_date="2026-01-31",
            )
        ),
        user,
        db,
    )
    active = create_assignment(
        AssignmentCreate(
            **assignment_payload(
                replacement_member.id,
                project_code="PROJECT-ACTIVE",
                allocation_percentage="40",
            )
        ),
        user,
        db,
    )

    resigned = list_assignments("", None, None, None, user, db, member_resigned=True)
    not_resigned = list_assignments("", None, None, None, user, db, member_resigned=False)
    assert [item.id for item in resigned.items] == [departing.id]
    assert [item.id for item in not_resigned.items] == [active.id]

    with pytest.raises(HTTPException) as error:
        create_assignment(
            AssignmentCreate(
                **assignment_payload(
                    departing_member.id,
                    project_code="PROJECT-OVERLAP",
                    allocation_percentage="41",
                    start_date="2026-01-31",
                    end_date="2026-02-01",
                )
            ),
            user,
            db,
        )
    assert error.value.status_code == 409

    shortened = update_assignment(
        departing.id,
        AssignmentUpdate(end_date=date(2026, 1, 30)),
        user,
        db,
    )
    assert shortened.end_date == date(2026, 1, 30)
    after_departure = create_assignment(
        AssignmentCreate(
            **assignment_payload(
                departing_member.id,
                project_code="PROJECT-AFTER-DEPARTURE",
                allocation_percentage="41",
                start_date="2026-01-31",
                end_date="2026-02-01",
            )
        ),
        user,
        db,
    )
    assert after_departure.member_id == departing_member.id

    replacement = create_assignment(
        AssignmentCreate(
            **assignment_payload(
                replacement_member.id,
                project_code="PROJECT-REPLACEMENT",
                allocation_percentage="60",
                start_date="2026-02-01",
                end_date="2026-02-28",
            )
        ),
        user,
        db,
    )
    assert replacement.member_id == replacement_member.id
    assert replacement.id != departing.id
    assert replacement.member_resigned is False


def test_export_assignments_returns_filtered_xlsx_with_ordered_columns(db: Session) -> None:
    user = CurrentUser(user_id="lead", role="Chapter Lead")
    first = create_test_member(
        db,
        user,
        dni="12345678",
        first_name="Ana",
        paternal_surname="Perez",
        maternal_surname=None,
        email="ana.export@example.com",
    )
    second = create_test_member(
        db,
        user,
        dni="87654321",
        first_name="Luis",
        paternal_surname="Ruiz",
        maternal_surname=None,
        email="luis.export@example.com",
    )
    planilla = create_test_member(
        db,
        user,
        dni="11223344",
        first_name="Marta",
        paternal_surname="Diaz",
        maternal_surname=None,
        email="marta.export@example.com",
        employment_type="PLANILLA",
        vendor_id=None,
        professional_role_id="role-2",
    )
    create_assignment(
        AssignmentCreate(
            **assignment_payload(
                first.id,
                project_code="PROJECT-ANA",
                start_date="2026-03-31",
                end_date="2026-04-30",
            )
        ),
        user,
        db,
    )
    create_assignment(
        AssignmentCreate(
            **assignment_payload(
                second.id,
                project_code="PROJECT-LUIS",
                start_date="2026-03-31",
                end_date="2026-05-31",
            )
        ),
        user,
        db,
    )
    create_assignment(
        AssignmentCreate(
            **assignment_payload(
                planilla.id,
                assigned_squad_id="squad-2",
                executor_squad_id="squad-2",
                project_code="PROJECT-MARTA",
                start_date="2026-01-01",
                end_date="2026-01-31",
            )
        ),
        user,
        db,
    )

    paged = list_assignments(
        "",
        "vendor-1",
        None,
        None,
        user,
        db,
        page_size=1,
        start_date=date(2026, 3, 31),
        end_date=date(2026, 3, 31),
    )
    assert paged.total == 2

    response = export_assignments(
        search="",
        vendor_id="vendor-1",
        professional_role_id=None,
        assigned_squad_id=None,
        start_date=date(2026, 3, 31),
        end_date=date(2026, 3, 31),
        sort_by="member_dni",
        sort_direction="asc",
        _=user,
        db=db,
    )
    workbook = load_workbook(BytesIO(response.body), data_only=True)
    worksheet = workbook["Asignaciones"]
    rows = list(worksheet.values)

    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert response.headers["content-disposition"] == 'attachment; filename="assignments.xlsx"'
    assert rows[0] == (
        "DNI",
        "Nombre Completo",
        "Proveedor",
        "Rol",
        "Squad Asignado",
        "Proyecto",
        "Fecha Inicio",
        "Fecha Fin",
        "% asignado",
    )
    normalized_rows = [
        row[:6] + (row[6].date(), row[7].date(), row[8])
        for row in rows[1:]
    ]
    assert normalized_rows == [
        ("12345678", "Ana Perez", "Proveedor Uno", "Data Engineer", "Squad Data", "PROJECT-ANA", date(2026, 3, 31), date(2026, 4, 30), 60),
        ("87654321", "Luis Ruiz", "Proveedor Uno", "Data Engineer", "Squad Data", "PROJECT-LUIS", date(2026, 3, 31), date(2026, 5, 31), 60),
    ]
    assert worksheet["G2"].number_format == "yyyy-mm-dd"
    assert worksheet["I2"].number_format == '0.##"%"'

    all_rows = list(
        load_workbook(
            BytesIO(
                export_assignments(
                    search="",
                    vendor_id=None,
                    professional_role_id=None,
                    assigned_squad_id=None,
                    _=user,
                    db=db,
                ).body
            ),
            data_only=True,
        )["Asignaciones"].values
    )
    assert any(row[0] == "11223344" and row[2] == "Planilla" for row in all_rows[1:])


def test_export_assignments_returns_empty_workbook_for_no_matches_and_validates_dates(db: Session) -> None:
    user = CurrentUser(user_id="lead", role="Chapter Lead")
    response = export_assignments(
        search="does-not-exist",
        vendor_id=None,
        professional_role_id=None,
        assigned_squad_id=None,
        _=user,
        db=db,
    )
    worksheet = load_workbook(BytesIO(response.body), data_only=True)["Asignaciones"]
    assert worksheet.max_row == 1

    with pytest.raises(HTTPException, match="filtro"):
        export_assignments(
            search="",
            vendor_id=None,
            professional_role_id=None,
            assigned_squad_id=None,
            start_date=date(2026, 2, 1),
            end_date=date(2026, 1, 1),
            _=user,
            db=db,
        )
